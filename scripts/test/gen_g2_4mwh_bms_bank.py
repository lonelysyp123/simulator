#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 G2 4MWh 标准版 BMS-EMS Modbus IO List 的 Bank Information 工作表生成 bms_bank.csv。

⚠ 本脚本全量重写目标 CSV。点表允许手工维护；日常增删请直接编辑 CSV。
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "【G2_4Mwh标准版】BMS-EMS Modbus IO List-V1.1- 2025-05-26.xlsx"
OUT_DIR = ROOT / "pointmaps" / "models" / "bms" / "g2_4mwh"
OUT_BANK = OUT_DIR / "bms_bank.csv"
OUT_RACK = OUT_DIR / "bms_rack.csv"

SKIP_NAME_RE = re.compile(r"(IP地址|子网掩码|端口号|port number)", re.I)
RANGE_ADDR_RE = re.compile(r"-")
HEX_ADDR_RE = re.compile(r"0[xX]([0-9A-Fa-f]+)")

SCALE_BY_UNIT = {
    "-": 1, "": 1, "kwh": 1, "kw": 1, "kω": 1, "ah": 1, "mω": 1,
    "0.1v": 10, "0.1a": 10, "0.1℃": 10, "0.1°c": 10,
    "‰": 1000, "mv": 1000,
}


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", str(s)).replace("\n", " ").strip()


def parse_bit(defn):
    m = re.match(r"Bit\s*(\d+)\s*[:：]?\s*(.+)", clean(defn), re.I)
    if not m:
        return None
    text = clean(m.group(2)).rstrip(";；")
    if text in ("…", "...", "……"):
        return None
    return int(m.group(1)), text


def parse_addr(raw):
    if raw is None:
        return None
    text = clean(raw)
    if not text or RANGE_ADDR_RE.search(text):
        return None
    m = HEX_ADDR_RE.search(text)
    if not m:
        return None
    return int(m.group(1), 16)


def scale_of(unit):
    key = clean(unit).lower().replace("℃", "℃")
    if key in SCALE_BY_UNIT:
        return SCALE_BY_UNIT[key]
    key2 = key.replace("°c", "℃")
    return SCALE_BY_UNIT.get(key2, 1)


def csv_type(xl_type, signed=False):
    t = clean(xl_type).lower()
    if t in ("int16", "s16"):
        return "int16"
    if t in ("int32", "s32"):
        return "int32"
    if t in ("u32", "uint32"):
        return "u32"
    if t in ("u16", "uint16", ""):
        return "int16" if signed else "u16"
    return "u16"


BINDINGS = {
    "一键并网开关": "BatteryStacks[0].GridConnectCommand",
    "复位": "BatteryStacks[0].FaultClearCommand",
    "一键并网开关状态": "BatteryStacks[0].GridConnectStatus",
    "系统运行状态": "BatteryStacks[0].OperationStatus",
    "系统充放电状态": "BatteryStacks[0].ChargeDischargeStatus",
    "系统三级报警汇总（轻微）": "BatteryStacks[0].BMSFaultSummary",
    "系统二级报警汇总": "BatteryStacks[0].BMSAlarmSummary",
    "系统一级报警汇总（严重）": "BatteryStacks[0].BMSProtectionSummary",
    "系统总电压": "BatteryStacks[0].TotalVoltage",
    "系统总电流": "BatteryStacks[0].Current",
    "系统soc": "BatteryStacks[0].SOC",
    "系统soh": "BatteryStacks[0].SOH",
    "系统绝缘值": "BatteryStacks[0].InsulationPlus",
    "系统可充电量": "BatteryStacks[0].AvailableChargeCapacity",
    "系统可放电量": "BatteryStacks[0].AvailableDischargeCapacity",
    "系统允许最大充电电流": "BatteryStacks[0].MaxChargeCurrent",
    "系统允许最大放电电流": "BatteryStacks[0].MaxDischargeCurrent",
    "簇间电流差异值": "BatteryStacks[0].ClusterCurrentDiff",
    "簇间总压差异值": "BatteryStacks[0].ClusterVoltageDiff",
    "系统最高电压电池所在rack号": "BatteryStacks[0].MaxCellVoltageClusterId",
    "系统最高电压电池所在组": "BatteryStacks[0].MaxCellVoltagePackId",
    "系统最高电压电池所在位置": "BatteryStacks[0].MaxCellVoltageCellId",
    "系统最高电池电压": "BatteryStacks[0].MaxCellVoltage",
    "系统最低电压电池所在rack号": "BatteryStacks[0].MinCellVoltageClusterId",
    "系统最低电压电池所在组": "BatteryStacks[0].MinCellVoltagePackId",
    "系统最低电压电池所在位置": "BatteryStacks[0].MinCellVoltageCellId",
    "系统最低电池电压": "BatteryStacks[0].MinCellVoltage",
    "系统平均电压": "BatteryStacks[0].AvgCellVoltage",
    "系统最高温度电池所在rack号": "BatteryStacks[0].MaxCellTempClusterId",
    "系统最高温度电池所在组": "BatteryStacks[0].MaxCellTempPackId",
    "系统最高温度电池所在位置": "BatteryStacks[0].MaxCellTempCellId",
    "系统最高电池温度": "BatteryStacks[0].MaxCellTemp",
    "系统最低温度电池所在rack号": "BatteryStacks[0].MinCellTempClusterId",
    "系统最低温度电池所在组": "BatteryStacks[0].MinCellTempPackId",
    "系统最低温度电池所在位置": "BatteryStacks[0].MinCellTempCellId",
    "系统最低电池温度": "BatteryStacks[0].MinCellTemp",
    "系统平均温度": "BatteryStacks[0].AvgCellTemp",
    "累计充电电量": "BatteryStacks[0].CumulativeChargeEnergy",
    "累计放电电量": "BatteryStacks[0].CumulativeDischargeEnergy",
    "当前在网簇数": "BatteryStacks[0].OnlineClusterCount",
    "堆内总簇数": "BatteryStacks[0].TotalClusterCount",
    "最大允许放电功率": "BatteryStacks[0].MaxDischargePower",
    "最大允许充电功率": "BatteryStacks[0].MaxChargePower",
    "rack三级报警汇总2": "BatteryStacks[0].BMSFaultSummary2",
    "rack二级报警汇总2": "BatteryStacks[0].BMSAlarmSummary2",
    "rack一级报警汇总2": "BatteryStacks[0].BMSProtectionSummary2",
    "集装箱环境温度": "TempHumiditySensors[0].Temperature",
}

CONSTANTS = {
    "堆内总簇数": "12",
    "最小并机簇数": "12",
}


def bind(name, scale):
    key = name.lower()
    path = BINDINGS.get(name) or BINDINGS.get(key)
    if path:
        return f"model=4|arg1=bmsdeviceId.{path}|arg2=|arg3=|arg4={scale}"
    return CONSTANTS.get(name, "0")


def describe(name, bits, defn):
    parts = [name]
    brief = clean(defn)
    if brief and not bits:
        brief = brief.replace(",", "，")
        if len(brief) > 80:
            brief = brief[:80].rstrip() + "…"
        parts.append(brief)
    if bits:
        bit_txt = ";".join(f"Bit{b} {t}" for b, t in bits)
        parts.append(bit_txt.replace(",", "，"))
    return " ".join(p for p in parts if p)


def parse_sheet(ws):
    rows = []
    last_name = ""
    pending_bits = []
    current_idx = None

    def flush_bits():
        nonlocal pending_bits, current_idx
        if current_idx is not None and pending_bits:
            rows[current_idx]["bits"] = pending_bits
        pending_bits = []

    for r in range(1, ws.max_row + 1):
        name = clean(ws.cell(row=r, column=1).value)
        en = clean(ws.cell(row=r, column=2).value)
        defn = clean(ws.cell(row=r, column=3).value)
        attr = clean(ws.cell(row=r, column=10).value).upper().replace(" ", "")
        typ = clean(ws.cell(row=r, column=11).value)
        unit = clean(ws.cell(row=r, column=9).value)
        addr = parse_addr(ws.cell(row=r, column=8).value)
        raw_addr = clean(ws.cell(row=r, column=8).value)

        if name in ("名称", "Name") or name.startswith("Bank Information") or name.startswith("寄存器") or name.startswith("Register") or name.startswith("Modbus Addr"):
            continue
        if name.startswith("Rack Enable"):
            continue

        if addr is None:
            bit = parse_bit(defn)
            if bit:
                pending_bits.append(bit)
            continue

        if RANGE_ADDR_RE.search(raw_addr) or SKIP_NAME_RE.search(name) or SKIP_NAME_RE.search(en):
            flush_bits()
            current_idx = None
            continue

        if not name:
            if last_name.startswith("时间") and defn:
                name = "时间-" + defn.split()[0]
            else:
                continue
        else:
            last_name = name

        flush_bits()
        is_write = "W" in attr or (attr == "" and addr < 0x1000)
        same_row_bit = parse_bit(defn)
        rows.append({
            "name": name,
            "en": en,
            "defn": "" if same_row_bit else defn,
            "addr": addr,
            "unit": unit,
            "typ": typ,
            "write": is_write,
            "bits": [same_row_bit] if same_row_bit else [],
        })
        current_idx = len(rows) - 1
        pending_bits = list(rows[current_idx]["bits"])

    flush_bits()
    return rows


def merge_energy(rows):
    merged = []
    i = 0
    while i < len(rows):
        cur = rows[i]
        if i + 1 < len(rows) and cur["name"].endswith("高位") and rows[i + 1]["name"].endswith("低位") and rows[i + 1]["addr"] == cur["addr"] + 1:
            base = cur["name"].replace("高位", "").replace("低位", "").strip()
            cur = dict(cur)
            cur["name"] = base
            cur["typ"] = "u32"
            cur["size"] = 32
            cur["defn"] = ""
            merged.append(cur)
            i += 2
            continue
        cur = dict(cur)
        cur["size"] = 16
        merged.append(cur)
        i += 1
    return merged


def to_csv_rows(regs):
    out = [["FunctionCode", "Address", "Type", "Size", "ParamName", "Scale", "Description", "ModelSim"]]
    writes = [r for r in regs if r["write"]]
    reads = [r for r in regs if not r["write"]]
    writes.sort(key=lambda x: x["addr"])
    reads.sort(key=lambda x: x["addr"])
    n = 1
    for r in writes + reads:
        scale = scale_of(r["unit"])
        if r["size"] == 32:
            typ = "u32"
        else:
            signed = (not r["typ"]) and ("电流" in r["name"] or "温度" in r["name"]) and ("所在" not in r["name"])
            typ = csv_type(r["typ"], signed=signed)
        fc = "6" if r["write"] else "3"
        desc = describe(r["name"], r.get("bits") or [], r.get("defn") or "")
        out.append([
            fc, str(r["addr"]), typ, str(r["size"]), f"param{n}", str(scale), desc, bind(r["name"], scale),
        ])
        n += 1
    return out


def main():
    if not XLSX.exists():
        print(f"找不到协议文件: {XLSX}", file=sys.stderr)
        return 1
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb["Bank Information"]
    regs = merge_energy(parse_sheet(ws))
    wb.close()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    rows = to_csv_rows(regs)
    with OUT_BANK.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(rows)

    if not OUT_RACK.exists():
        with OUT_RACK.open("w", newline="", encoding="utf-8") as f:
            csv.writer(f).writerow(
                ["FunctionCode", "Address", "Type", "Size", "ParamName", "Scale", "Description", "ModelSim"]
            )

    (OUT_DIR / "model.json").write_text(
        '{\n'
        '  "id": "g2_4mwh",\n'
        '  "name": "G2 4MWh 标准版",\n'
        '  "description": "G2 4MWh 标准版 BMS-EMS IO List V1.1 Bank Information；rack 点表待 Rack Information 工作表导入"\n'
        '}\n',
        encoding="utf-8",
    )

    n_read = sum(1 for r in rows[1:] if r[0] == "3")
    n_write = sum(1 for r in rows[1:] if r[0] == "6")
    n_bind = sum(1 for r in rows[1:] if str(r[7]).startswith("model="))
    print(f"G2 4MWh bank: {len(rows) - 1} 点 (读 {n_read} / 写 {n_write}), 绑定 {n_bind} -> {OUT_BANK}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
